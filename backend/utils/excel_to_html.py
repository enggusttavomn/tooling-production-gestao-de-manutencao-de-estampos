"""
Arquivo: backend/utils/excel_to_html.py
Descricao: converte uma aba de planilha Excel em HTML preservando estilo visual.

Objetivo:
- Facilitar visualizacao web da planilha padrao sem depender do Excel no cliente.
- Manter formato aproximado da planilha (cores, fonte, bordas, merge e dimensoes).
"""

import argparse
import datetime as dt
import html
# Notas de manutencao:
# - Objetivo: manter comportamento funcional sem quebrar compatibilidade.
# - Cuidado: validar efeitos em fluxo web, banco e exportacoes.
# - Ao alterar: preferir mudancas pequenas e validacao local.

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Color
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries


DEFAULT_COL_WIDTH = 8.43  # Excel default column width
DEFAULT_ROW_HEIGHT_PT = 15  # Excel default row height in points


def cor_para_hex(cor: Color | None) -> str | None:
    """Converte cor do openpyxl para hexadecimal CSS quando possivel."""
    if cor is None:
        return None
    if cor.type == "rgb" and cor.rgb:
        rgb = cor.rgb[-6:]
        return f"#{rgb}"
    if cor.type == "indexed" and cor.indexed is not None:
        idx = int(cor.indexed)
        if 0 <= idx < len(COLOR_INDEX):
            rgb = COLOR_INDEX[idx]
            if rgb:
                return f"#{rgb[-6:]}"
    return None


def borda_para_css(lado) -> str | None:
    """Mapeia uma borda do Excel para declaracao CSS simples."""
    if lado is None or lado.style is None:
        return None
    cor = cor_para_hex(lado.color) or "#000000"
    return f"1px solid {cor}"


def valor_para_html(celula) -> str:
    """Converte valor da celula para string HTML segura."""
    valor = celula.value
    if valor is None:
        return "&nbsp;"
    if isinstance(valor, dt.datetime):
        return html.escape(valor.strftime("%d-%b-%Y"))
    if isinstance(valor, dt.date):
        return html.escape(valor.strftime("%d-%b-%Y"))
    if isinstance(valor, float) and is_date_format(celula.number_format):
        return html.escape(str(valor))
    texto = str(valor)
    texto = html.escape(texto)
    return texto.replace("\n", "<br>")


def montar_estilos(celula) -> str:
    """Monta estilo inline CSS aproximando o visual original da celula."""
    estilos = []

    fonte = celula.font
    if fonte:
        if fonte.name:
            estilos.append(f"font-family: {fonte.name}, Calibri, Arial, sans-serif;")
        if fonte.sz:
            estilos.append(f"font-size: {fonte.sz}px;")
        if fonte.bold:
            estilos.append("font-weight: bold;")
        if fonte.italic:
            estilos.append("font-style: italic;")
        if fonte.color:
            cor = cor_para_hex(fonte.color)
            if cor:
                estilos.append(f"color: {cor};")

    alinhamento = celula.alignment
    if alinhamento:
        if alinhamento.horizontal:
            estilos.append(f"text-align: {alinhamento.horizontal};")
        if alinhamento.vertical:
            estilos.append(f"vertical-align: {alinhamento.vertical};")
        if alinhamento.wrap_text:
            estilos.append("white-space: normal;")

    preenchimento = celula.fill
    if preenchimento and getattr(preenchimento, "patternType", None) == "solid":
        cor = cor_para_hex(preenchimento.fgColor)
        if cor:
            estilos.append(f"background-color: {cor};")

    borda = celula.border
    if borda:
        topo = borda_para_css(borda.top)
        direita = borda_para_css(borda.right)
        baixo = borda_para_css(borda.bottom)
        esquerda = borda_para_css(borda.left)
        if topo:
            estilos.append(f"border-top: {topo};")
        if direita:
            estilos.append(f"border-right: {direita};")
        if baixo:
            estilos.append(f"border-bottom: {baixo};")
        if esquerda:
            estilos.append(f"border-left: {esquerda};")

    if not any(e.startswith("padding") for e in estilos):
        estilos.append("padding: 2px 4px;")

    return " ".join(estilos)


def exportar_aba_para_html(caminho_entrada: Path, nome_aba: str, caminho_saida: Path) -> None:
    """Renderiza uma aba da planilha para HTML com colunas/linhas e merges."""
    pasta_trabalho = load_workbook(caminho_entrada, data_only=True)
    if nome_aba not in pasta_trabalho.sheetnames:
        raise ValueError(f"Aba '{nome_aba}' nao encontrada. Disponiveis: {', '.join(pasta_trabalho.sheetnames)}")
    planilha = pasta_trabalho[nome_aba]

    dimensao = planilha.calculate_dimension()
    col_min, lin_min, col_max, lin_max = range_boundaries(dimensao)

    # Estruturas auxiliares para respeitar merged cells no HTML final.
    mapa_mescladas: dict[tuple[int, int], tuple[int, int]] = {}
    celulas_ignoradas: set[tuple[int, int]] = set()
    for celula_mesclada in planilha.merged_cells.ranges:
        m_col_min, m_lin_min, m_col_max, m_lin_max = range_boundaries(str(celula_mesclada))
        mapa_mescladas[(m_lin_min, m_col_min)] = (m_lin_max - m_lin_min + 1, m_col_max - m_col_min + 1)
        for lin in range(m_lin_min, m_lin_max + 1):
            for col in range(m_col_min, m_col_max + 1):
                if (lin, col) != (m_lin_min, m_col_min):
                    celulas_ignoradas.add((lin, col))

    # Converte larguras/alturas do Excel para pixels aproximados.
    larguras_colunas = []
    for col in range(col_min, col_max + 1):
        letra = get_column_letter(col)
        dimensao_coluna = planilha.column_dimensions.get(letra)
        largura = dimensao_coluna.width if dimensao_coluna and dimensao_coluna.width else DEFAULT_COL_WIDTH
        pixels = int(largura * 7 + 5)
        larguras_colunas.append(pixels)

    alturas_linhas = []
    for lin in range(lin_min, lin_max + 1):
        dimensao_linha = planilha.row_dimensions.get(lin)
        altura = dimensao_linha.height if dimensao_linha and dimensao_linha.height else DEFAULT_ROW_HEIGHT_PT
        pixels = int(altura * 1.333)
        alturas_linhas.append(pixels)

    # Monta o documento HTML final em memoria e grava de uma vez no arquivo.
    linhas_html = []
    linhas_html.append("<!DOCTYPE html>")
    linhas_html.append("<html lang=\"pt-BR\">")
    linhas_html.append("<head>")
    linhas_html.append("<meta charset=\"utf-8\">")
    linhas_html.append("<meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\">")
    linhas_html.append(f"<title>{html.escape(nome_aba)}</title>")
    linhas_html.append("<style>")
    linhas_html.append("body { margin: 16px; font-family: Calibri, Arial, sans-serif; background: #f5f6f7; }")
    linhas_html.append(".sheet-wrap { background: #ffffff; padding: 16px; box-shadow: 0 2px 12px rgba(0,0,0,0.08); }")
    linhas_html.append("table { border-collapse: collapse; }")
    linhas_html.append("td { min-width: 6px; }")
    linhas_html.append("</style>")
    linhas_html.append("</head>")
    linhas_html.append("<body>")
    linhas_html.append("<div class=\"sheet-wrap\">")
    linhas_html.append("<table>")
    linhas_html.append("<colgroup>")
    for largura in larguras_colunas:
        linhas_html.append(f"<col style=\"width: {largura}px;\">")
    linhas_html.append("</colgroup>")

    for idx_linha, lin in enumerate(range(lin_min, lin_max + 1), start=0):
        altura = alturas_linhas[idx_linha]
        linhas_html.append(f"<tr style=\"height: {altura}px;\">")
        for col in range(col_min, col_max + 1):
            if (lin, col) in celulas_ignoradas:
                continue
            celula = planilha.cell(row=lin, column=col)
            rowspan, colspan = mapa_mescladas.get((lin, col), (1, 1))
            atributos_span = ""
            if rowspan > 1:
                atributos_span += f" rowspan=\"{rowspan}\""
            if colspan > 1:
                atributos_span += f" colspan=\"{colspan}\""
            estilos = montar_estilos(celula)
            valor = valor_para_html(celula)
            linhas_html.append(f"<td{atributos_span} style=\"{estilos}\">{valor}</td>")
        linhas_html.append("</tr>")

    linhas_html.append("</table>")
    linhas_html.append("</div>")
    linhas_html.append("</body>")
    linhas_html.append("</html>")

    caminho_saida.write_text("\n".join(linhas_html), encoding="utf-8")


def main() -> None:
    """Entrada CLI para uso em scripts de automacao."""
    parser = argparse.ArgumentParser(description="Exporta uma aba de Excel para HTML.")
    parser.add_argument("--input", required=True, help="Caminho do arquivo .xlsx")
    parser.add_argument("--sheet", required=True, help="Nome da aba para exportar")
    parser.add_argument("--output", required=True, help="Caminho do HTML de saida")
    args = parser.parse_args()

    caminho_entrada = Path(args.input)
    caminho_saida = Path(args.output)
    exportar_aba_para_html(caminho_entrada, args.sheet, caminho_saida)


if __name__ == "__main__":
    main()
