import argparse
import datetime as dt
import html
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Color
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries


DEFAULT_COL_WIDTH = 8.43  # Excel default column width
DEFAULT_ROW_HEIGHT_PT = 15  # Excel default row height in points


def color_to_hex(color: Color | None) -> str | None:
    if color is None:
        return None
    if color.type == "rgb" and color.rgb:
        rgb = color.rgb[-6:]
        return f"#{rgb}"
    if color.type == "indexed" and color.indexed is not None:
        idx = int(color.indexed)
        if 0 <= idx < len(COLOR_INDEX):
            rgb = COLOR_INDEX[idx]
            if rgb:
                return f"#{rgb[-6:]}"
    return None


def border_to_css(side) -> str | None:
    if side is None or side.style is None:
        return None
    color = color_to_hex(side.color) or "#000000"
    return f"1px solid {color}"


def value_to_html(cell) -> str:
    value = cell.value
    if value is None:
        return "&nbsp;"
    if isinstance(value, dt.datetime):
        return html.escape(value.strftime("%d-%b-%Y"))
    if isinstance(value, dt.date):
        return html.escape(value.strftime("%d-%b-%Y"))
    if isinstance(value, float) and is_date_format(cell.number_format):
        return html.escape(str(value))
    text = str(value)
    text = html.escape(text)
    return text.replace("\n", "<br>")


def build_styles(cell) -> str:
    styles = []

    font = cell.font
    if font:
        if font.name:
            styles.append(f"font-family: {font.name}, Calibri, Arial, sans-serif;")
        if font.sz:
            styles.append(f"font-size: {font.sz}px;")
        if font.bold:
            styles.append("font-weight: bold;")
        if font.italic:
            styles.append("font-style: italic;")
        if font.color:
            color = color_to_hex(font.color)
            if color:
                styles.append(f"color: {color};")

    alignment = cell.alignment
    if alignment:
        if alignment.horizontal:
            styles.append(f"text-align: {alignment.horizontal};")
        if alignment.vertical:
            styles.append(f"vertical-align: {alignment.vertical};")
        if alignment.wrap_text:
            styles.append("white-space: normal;")

    fill = cell.fill
    if fill and getattr(fill, "patternType", None) == "solid":
        color = color_to_hex(fill.fgColor)
        if color:
            styles.append(f"background-color: {color};")

    border = cell.border
    if border:
        top = border_to_css(border.top)
        right = border_to_css(border.right)
        bottom = border_to_css(border.bottom)
        left = border_to_css(border.left)
        if top:
            styles.append(f"border-top: {top};")
        if right:
            styles.append(f"border-right: {right};")
        if bottom:
            styles.append(f"border-bottom: {bottom};")
        if left:
            styles.append(f"border-left: {left};")

    if not any(s.startswith("padding") for s in styles):
        styles.append("padding: 2px 4px;")

    return " ".join(styles)


def export_sheet_to_html(input_path: Path, sheet_name: str, output_path: Path) -> None:
    wb = load_workbook(input_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {', '.join(wb.sheetnames)}")
    ws = wb[sheet_name]

    dimension = ws.calculate_dimension()
    min_col, min_row, max_col, max_row = range_boundaries(dimension)

    merged_map: dict[tuple[int, int], tuple[int, int]] = {}
    merged_skip: set[tuple[int, int]] = set()
    for merged in ws.merged_cells.ranges:
        m_min_col, m_min_row, m_max_col, m_max_row = range_boundaries(str(merged))
        merged_map[(m_min_row, m_min_col)] = (m_max_row - m_min_row + 1, m_max_col - m_min_col + 1)
        for r in range(m_min_row, m_max_row + 1):
            for c in range(m_min_col, m_max_col + 1):
                if (r, c) != (m_min_row, m_min_col):
                    merged_skip.add((r, c))

    col_widths = []
    for col in range(min_col, max_col + 1):
        letter = get_column_letter(col)
        dim = ws.column_dimensions.get(letter)
        width = dim.width if dim and dim.width else DEFAULT_COL_WIDTH
        px = int(width * 7 + 5)
        col_widths.append(px)

    row_heights = []
    for row in range(min_row, max_row + 1):
        dim = ws.row_dimensions.get(row)
        height = dim.height if dim and dim.height else DEFAULT_ROW_HEIGHT_PT
        px = int(height * 1.333)
        row_heights.append(px)

    html_lines = []
    html_lines.append("<!DOCTYPE html>")
    html_lines.append("<html lang=\"pt-BR\">")
    html_lines.append("<head>")
    html_lines.append("<meta charset=\"utf-8\">")
    html_lines.append("<meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\">")
    html_lines.append(f"<title>{html.escape(sheet_name)}</title>")
    html_lines.append("<style>")
    html_lines.append("body { margin: 16px; font-family: Calibri, Arial, sans-serif; background: #f5f6f7; }")
    html_lines.append(".sheet-wrap { background: #ffffff; padding: 16px; box-shadow: 0 2px 12px rgba(0,0,0,0.08); }")
    html_lines.append("table { border-collapse: collapse; }")
    html_lines.append("td { min-width: 6px; }")
    html_lines.append("</style>")
    html_lines.append("</head>")
    html_lines.append("<body>")
    html_lines.append("<div class=\"sheet-wrap\">")
    html_lines.append("<table>")
    html_lines.append("<colgroup>")
    for width in col_widths:
        html_lines.append(f"<col style=\"width: {width}px;\">")
    html_lines.append("</colgroup>")

    for row_idx, row in enumerate(range(min_row, max_row + 1), start=0):
        height = row_heights[row_idx]
        html_lines.append(f"<tr style=\"height: {height}px;\">")
        for col in range(min_col, max_col + 1):
            if (row, col) in merged_skip:
                continue
            cell = ws.cell(row=row, column=col)
            rowspan, colspan = merged_map.get((row, col), (1, 1))
            span_attrs = ""
            if rowspan > 1:
                span_attrs += f" rowspan=\"{rowspan}\""
            if colspan > 1:
                span_attrs += f" colspan=\"{colspan}\""
            styles = build_styles(cell)
            value = value_to_html(cell)
            html_lines.append(f"<td{span_attrs} style=\"{styles}\">{value}</td>")
        html_lines.append("</tr>")

    html_lines.append("</table>")
    html_lines.append("</div>")
    html_lines.append("</body>")
    html_lines.append("</html>")

    output_path.write_text("\n".join(html_lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Export Excel sheet to HTML.")
    parser.add_argument("--input", required=True, help="Path to the .xlsx file")
    parser.add_argument("--sheet", required=True, help="Sheet name to export")
    parser.add_argument("--output", required=True, help="Path to output HTML")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)
    export_sheet_to_html(input_path, args.sheet, output_path)


if __name__ == "__main__":
    main()
