from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import json
import os
import subprocess

from flask import Flask, jsonify, make_response, request, send_file, send_from_directory
from flask_cors import CORS
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

try:
    import mysql.connector
    from mysql.connector import Error as MySQLError
except ImportError:
    mysql = None
    MySQLError = Exception

try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

if load_dotenv:
    load_dotenv()
else:
    env_path = os.path.join(os.path.dirname(__file__), ".env")
    if os.path.exists(env_path):
        with open(env_path, "r", encoding="utf-8") as env_file:
            for raw_line in env_file:
                line = raw_line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                key, value = line.split("=", 1)
                key = key.strip()
                value = value.strip().strip('"').strip("'")
                if key and key not in os.environ:
                    os.environ[key] = value


app = Flask(__name__, static_folder=None)
CORS(app)

DEV_RELOAD_SNIPPET = """
<script>
(function () {
  if (window.location.hostname !== "localhost" && window.location.hostname !== "127.0.0.1") {
    return;
  }
  let lastTs = 0;

  async function checkReload() {
    try {
      const response = await fetch("/__dev_timestamp?ts=" + Date.now(), { cache: "no-store" });
      if (!response.ok) return;
      const data = await response.json();
      const currentTs = Number(data.ts || 0);
      if (!currentTs) return;
      if (lastTs && currentTs > lastTs) {
        window.location.reload();
        return;
      }
      lastTs = currentTs;
    } catch (_error) {
      // Silent in dev polling.
    }
  }

  checkReload();
  setInterval(checkReload, 1000);
})();
</script>
"""


def _serve_html_with_dev_reload(file_path):
    try:
        with open(file_path, "r", encoding="utf-8") as file:
            html = file.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as file:
            html = file.read()

    if "</body>" in html:
        html = html.replace("</body>", DEV_RELOAD_SNIPPET + "\n</body>", 1)
    else:
        html = html + DEV_RELOAD_SNIPPET

    response = make_response(html)
    response.headers["Content-Type"] = "text/html; charset=utf-8"
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


def _db_config():
    config = {
        "host": os.getenv("DB_HOST", "localhost"),
        "port": int(os.getenv("DB_PORT", "3306")),
        "user": os.getenv("DB_USER", ""),
        "password": os.getenv("DB_PASSWORD", ""),
        "database": os.getenv("DB_NAME", ""),
        "charset": "utf8mb4",
        "use_unicode": True,
        "autocommit": False,
    }

    ssl_mode = os.getenv("DB_SSL_MODE", "").strip().upper()
    if ssl_mode == "DISABLED":
        config["ssl_disabled"] = True
    if ssl_mode == "REQUIRED":
        config["ssl_disabled"] = False

    return config


def _connect_db():
    if mysql is None:
        raise RuntimeError("Dependencia mysql-connector-python nao instalada.")

    cfg = _db_config()
    missing = [key for key in ("DB_USER", "DB_PASSWORD", "DB_NAME") if not os.getenv(key)]
    if missing:
        raise RuntimeError("Variaveis ausentes: " + ", ".join(missing))

    return mysql.connector.connect(**cfg)


def _escape_sql_value(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    text = str(value)
    text = text.replace("\\", "\\\\").replace("'", "\\'")
    return f"'{text}'"


def _interpolate_query(query, params=None):
    if not params:
        return query
    parts = query.split("%s")
    if len(parts) - 1 != len(params):
        raise RuntimeError("Quantidade de parametros nao confere com placeholders SQL.")
    output = [parts[0]]
    for idx, value in enumerate(params):
        output.append(_escape_sql_value(value))
        output.append(parts[idx + 1])
    return "".join(output)


def _query_all_cli(query, params=None):
    cfg = _db_config()
    missing = [key for key in ("DB_HOST", "DB_PORT", "DB_USER", "DB_PASSWORD", "DB_NAME") if not os.getenv(key)]
    if missing:
        raise RuntimeError("Variaveis ausentes: " + ", ".join(missing))

    client_exe = os.getenv("DB_CLIENT_EXE", r"C:\Program Files\MariaDB 12.2\bin\mariadb.exe")
    if not os.path.exists(client_exe):
        raise RuntimeError(f"Cliente MariaDB nao encontrado em: {client_exe}")

    sql = _interpolate_query(query, params)
    args = [
        client_exe,
        "--skip-ssl",
        "--batch",
        "-h",
        str(cfg["host"]),
        "-P",
        str(cfg["port"]),
        "-u",
        str(cfg["user"]),
        f"-p{cfg['password']}",
        str(cfg["database"]),
        "-e",
        sql,
    ]

    completed = subprocess.run(
        args,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if completed.returncode != 0:
        stderr = (completed.stderr or "").strip()
        raise RuntimeError(stderr or "Falha ao executar consulta via cliente MariaDB.")

    lines = [line for line in (completed.stdout or "").splitlines() if line.strip() != ""]
    if not lines:
        return []

    headers = lines[0].split("\t")
    rows = []
    for line in lines[1:]:
        cols = line.split("\t")
        row = {}
        for idx, header in enumerate(headers):
            row[header] = cols[idx] if idx < len(cols) else None
        rows.append(_serialize_row(row))
    return rows


def _serialize_value(value):
    if isinstance(value, str):
        return _fix_mojibake(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _fix_mojibake(text):
    # Heuristic repair for common UTF-8/latin1 mojibake such as
    # "DISPOSITIVO ... ESPAÃƒâ€¡ADORES" or "LiberaÃƒÆ’Ã‚Â§ÃƒÆ’Ã‚Â£o".
    if not isinstance(text, str):
        return text
    if "Ãƒ" not in text and "Ã‚" not in text and "Ã¢" not in text:
        return text

    repaired = text
    for _ in range(3):
        try:
            candidate = repaired.encode("latin1", errors="strict").decode("utf-8", errors="strict")
        except Exception:
            break
        if candidate == repaired:
            break
        repaired = candidate
        if "Ãƒ" not in repaired and "Ã‚" not in repaired and "Ã¢" not in repaired:
            break
    return repaired


def _serialize_row(row):
    return {key: _serialize_value(value) for key, value in row.items()}


def _normalize_status(status_roteiro):
    value = (status_roteiro or "").strip().lower()
    if not value or value == "na":
        return "pendente"
    if "aprovado" in value or "conclu" in value:
        return "concluida"
    if "andamento" in value or "execu" in value:
        return "em_andamento"
    if "roteiro" in value:
        return "pendente"
    return "em_andamento"


def _map_banco_dados_to_api(row):
    status = _normalize_status(row.get("status_roteiro"))
    return {
        "id": row.get("numero_ficha"),
        "numero_ficha": row.get("numero_ficha"),
        "equipamento": row.get("tipo_lamina_processo"),
        "tag": row.get("numero_ferramental"),
        "tipo_manutencao": row.get("tipo_lamina_processo"),
        "data_execucao": _serialize_value(row.get("data_termino_producao")),
        "responsavel": row.get("operador_responsavel"),
        "area": row.get("pt"),
        "descricao": row.get("descricao_ferramental"),
        "observacoes": row.get("observacao_ultima_manutencao_ferramenta") or row.get("observacao_rnc"),
        "tempo_execucao": row.get("vida_util_ferramenta"),
        "status": status,
        "data_criacao": _serialize_value(row.get("data_emissao_requisicao")),
        "data_atualizacao": _serialize_value(row.get("data_termino_producao")),
        # Campos esperados pelo frontend atual:
        "descricao_ferramenta": row.get("descricao_ferramental"),
        "codigo_ferramenta": row.get("numero_ferramental"),
        "documento": row.get("numero_plan_interno") or row.get("numero_po"),
        "data_cadastro": _serialize_value(row.get("data_emissao_requisicao")),
        "responsavel_abertura": row.get("operador_responsavel"),
        "ferramenteiro": row.get("ferramenteiro_responsavel"),
        "qtde_afiacao": row.get("qtde_afiacao"),
        "data_termino": _serialize_value(row.get("data_termino_producao")),
        "descricao_trabalhos": row.get("tipo_lamina_processo"),
        "job": row.get("job"),
        "job_envolvidas": row.get("job_envolvidas"),
        "numero_ferramental": row.get("numero_ferramental"),
        "numero_desenho_ferramental": row.get("numero_desenho_ferramental"),
        "status_roteiro": row.get("status_roteiro"),
        "status_ferramental": row.get("status_ferramental"),
        "status_entrega_ferramental": row.get("status_entrega_ferramental"),
        "data_entrega_evento": _serialize_value(row.get("data_entrega_evento")),
        "data_termino_producao": _serialize_value(row.get("data_termino_producao")),
        "prensa": row.get("prensa"),
        "numero_rnc": row.get("numero_rnc"),
        "observacao_rnc": row.get("observacao_rnc"),
        "total_golpes_final": row.get("total_golpes_final"),
        "total_golpes_sem_afiar": row.get("total_golpes_sem_afiar"),
        "golpes_afiacao_1": row.get("1_golpes_afiacao"),
        "golpes_afiacao_2": row.get("2_golpes_afiacao"),
        "golpes_afiacao_3": row.get("3_golpes_afiacao"),
        "golpes_afiacao_4": row.get("4_golpes_afiacao"),
        "golpes_afiacao_5": row.get("5_golpes_afiacao"),
        "qtde_total_afiacao": row.get("qtde_total_afiacao"),
        "vida_util_ferramenta": row.get("vida_util_ferramenta"),
        "data_envio_ferramentaria": _serialize_value(row.get("data_envio_ferramentaria")),
        "horario_inicio_ferramenteiro": _serialize_value(row.get("horario_inicio_ferramenteiro")),
        "horario_termino_ferramenteiro": _serialize_value(row.get("horario_termino_ferramenteiro")),
    }


def _request_data():
    if request.is_json:
        return request.get_json(silent=True) or {}
    return request.form.to_dict(flat=True)


def _query_all(query, params=None):
    if mysql is None:
        return _query_all_cli(query, params)

    connection = _connect_db()
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute(query, params or ())
        rows = cursor.fetchall()
        return [_serialize_row(row) for row in rows]
    finally:
        cursor.close()
        connection.close()


def _query_one(query, params=None):
    if mysql is None:
        rows = _query_all_cli(query, params)
        return rows[0] if rows else None

    connection = _connect_db()
    cursor = connection.cursor(dictionary=True)
    try:
        cursor.execute(query, params or ())
        row = cursor.fetchone()
        return _serialize_row(row) if row else None
    finally:
        cursor.close()
        connection.close()


def _execute_write(query, params=None):
    connection = _connect_db()
    cursor = connection.cursor()
    try:
        cursor.execute(query, params or ())
        connection.commit()
        return cursor.lastrowid, cursor.rowcount
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()
        connection.close()


@app.route("/")
def home():
    return _serve_html_with_dev_reload(os.path.join("frontend", "login.html"))


@app.route("/<path:path>")
def static_files(path):
    if path.lower().endswith(".html") and os.path.isfile(path):
        return _serve_html_with_dev_reload(path)
    return send_from_directory(".", path)


@app.route("/api/fichas", methods=["POST"])
def salvar_ficha():
    try:
        return jsonify(
            {
                "success": False,
                "message": "API em modo somente leitura. Criacao de ficha desabilitada.",
            }
        ), 405
    except (MySQLError, RuntimeError, ValueError) as error:
        return jsonify({"success": False, "message": str(error)}), 400


@app.route("/api/fichas", methods=["GET"])
def buscar_fichas():
    try:
        search = request.args.get("search", "").strip().lower()
        tipo = request.args.get("tipo", "").strip()
        status = request.args.get("status", "").strip()
        data_inicio = request.args.get("data_inicio", "").strip()
        data_fim = request.args.get("data_fim", "").strip()

        query = """
            SELECT
                numero_ficha,
                REPLACE(REPLACE(REPLACE(COALESCE(job, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS job,
                REPLACE(REPLACE(REPLACE(COALESCE(job_envolvidas, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS job_envolvidas,
                REPLACE(REPLACE(REPLACE(COALESCE(tipo_lamina_processo, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS tipo_lamina_processo,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_desenho_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_desenho_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(status_roteiro, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS status_roteiro,
                REPLACE(REPLACE(REPLACE(COALESCE(status_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS status_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(status_entrega_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS status_entrega_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(descricao_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS descricao_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(pt, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS pt,
                data_entrega_evento, data_emissao_requisicao,
                REPLACE(REPLACE(REPLACE(COALESCE(operador_responsavel, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS operador_responsavel,
                REPLACE(REPLACE(REPLACE(COALESCE(ferramenteiro_responsavel, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS ferramenteiro_responsavel,
                qtde_afiacao,
                data_termino_producao,
                REPLACE(REPLACE(REPLACE(COALESCE(observacao_ultima_manutencao_ferramenta, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS observacao_ultima_manutencao_ferramenta,
                REPLACE(REPLACE(REPLACE(COALESCE(observacao_rnc, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS observacao_rnc,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_plan_interno, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_plan_interno,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_po, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_po,
                vida_util_ferramenta,
                REPLACE(REPLACE(REPLACE(COALESCE(prensa, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS prensa,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_rnc, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_rnc,
                total_golpes_final, total_golpes_sem_afiar,
                `1_golpes_afiacao`, `2_golpes_afiacao`, `3_golpes_afiacao`, `4_golpes_afiacao`, `5_golpes_afiacao`,
                qtde_total_afiacao, data_envio_ferramentaria, horario_inicio_ferramenteiro, horario_termino_ferramenteiro
            FROM banco_dados
            WHERE 1=1
        """
        params = []

        if search:
            query += " AND (LOWER(COALESCE(job, '')) LIKE %s OR LOWER(COALESCE(job_envolvidas, '')) LIKE %s OR LOWER(COALESCE(numero_ferramental, '')) LIKE %s)"
            like_value = f"%{search}%"
            params.extend([like_value, like_value, like_value])
        if tipo:
            query += " AND tipo_lamina_processo = %s"
            params.append(tipo)
        if status:
            query += " AND status_roteiro = %s"
            params.append(status)
        if data_inicio:
            query += " AND data_termino_producao >= %s"
            params.append(data_inicio)
        if data_fim:
            query += " AND data_termino_producao <= %s"
            params.append(data_fim)

        query += " ORDER BY numero_ficha DESC"
        rows = _query_all(query, params)
        return jsonify([_map_banco_dados_to_api(row) for row in rows]), 200
    except (MySQLError, RuntimeError) as error:
        return jsonify({"error": str(error)}), 400


@app.route("/api/estatisticas", methods=["GET"])
def obter_estatisticas():
    try:
        query = """
            SELECT
                COUNT(*) AS total,
                SUM(CASE WHEN LOWER(COALESCE(status_roteiro, '')) LIKE '%aprovado%' OR LOWER(COALESCE(status_roteiro, '')) LIKE '%conclu%' THEN 1 ELSE 0 END) AS concluidas,
                SUM(CASE WHEN LOWER(COALESCE(status_roteiro, '')) LIKE '%roteiro%' OR status_roteiro IS NULL OR LOWER(status_roteiro) = 'na' THEN 1 ELSE 0 END) AS pendentes,
                SUM(CASE WHEN NOT (LOWER(COALESCE(status_roteiro, '')) LIKE '%aprovado%' OR LOWER(COALESCE(status_roteiro, '')) LIKE '%conclu%' OR LOWER(COALESCE(status_roteiro, '')) LIKE '%roteiro%' OR status_roteiro IS NULL OR LOWER(status_roteiro) = 'na') THEN 1 ELSE 0 END) AS em_andamento,
                COALESCE(SUM(vida_util_ferramenta), 0) AS tempo_total
            FROM banco_dados
        """
        row = _query_one(query) or {}
        stats = {
            "total": int(row.get("total", 0) or 0),
            "pendentes": int(row.get("pendentes", 0) or 0),
            "em_andamento": int(row.get("em_andamento", 0) or 0),
            "concluidas": int(row.get("concluidas", 0) or 0),
            "tempo_total": float(row.get("tempo_total", 0) or 0),
        }
        return jsonify(stats), 200
    except (MySQLError, RuntimeError) as error:
        return jsonify({"error": str(error)}), 400


@app.route("/api/exportar-excel", methods=["GET"])
def exportar_excel():
    try:
        fichas = _query_all(
            """
            SELECT
                numero_ficha, job, job_envolvidas, tipo_lamina_processo, numero_ferramental,
                status_roteiro, descricao_ferramental, pt, data_emissao_requisicao,
                operador_responsavel, ferramenteiro_responsavel, qtde_afiacao,
                data_termino_producao, observacao_ultima_manutencao_ferramenta
            FROM banco_dados
            ORDER BY numero_ficha DESC
            """
        )

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Fichas de Manutencao"

        header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        headers = [
            "Numero Ficha",
            "Job",
            "Jobs Envolvidas",
            "Tipo Lamina Processo",
            "Numero Ferramental",
            "Status Roteiro",
            "Descricao Ferramental",
            "PT",
            "Data Emissao Requisicao",
            "Operador Responsavel",
            "Ferramenteiro Responsavel",
            "Qtde Afiacao",
            "Data Termino Producao",
            "Observacao Ultima Manutencao",
        ]

        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row_index, ficha in enumerate(fichas, 2):
            sheet.cell(row=row_index, column=1, value=ficha.get("numero_ficha"))
            sheet.cell(row=row_index, column=2, value=ficha.get("job"))
            sheet.cell(row=row_index, column=3, value=ficha.get("job_envolvidas"))
            sheet.cell(row=row_index, column=4, value=ficha.get("tipo_lamina_processo"))
            sheet.cell(row=row_index, column=5, value=ficha.get("numero_ferramental"))
            sheet.cell(row=row_index, column=6, value=ficha.get("status_roteiro"))
            sheet.cell(row=row_index, column=7, value=ficha.get("descricao_ferramental"))
            sheet.cell(row=row_index, column=8, value=ficha.get("pt"))
            sheet.cell(row=row_index, column=9, value=ficha.get("data_emissao_requisicao"))
            sheet.cell(row=row_index, column=10, value=ficha.get("operador_responsavel"))
            sheet.cell(row=row_index, column=11, value=ficha.get("ferramenteiro_responsavel"))
            sheet.cell(row=row_index, column=12, value=ficha.get("qtde_afiacao"))
            sheet.cell(row=row_index, column=13, value=ficha.get("data_termino_producao"))
            sheet.cell(row=row_index, column=14, value=ficha.get("observacao_ultima_manutencao_ferramenta"))

        for col in range(1, len(headers) + 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        filename = f"Fichas_Manutencao_{datetime.now().strftime('%d_%m_%Y')}.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )
    except (MySQLError, RuntimeError) as error:
        return jsonify({"error": str(error)}), 400


@app.route("/api/ficha/<int:ficha_id>", methods=["GET"])
def obter_ficha(ficha_id):
    try:
        ficha = _query_one(
            """
            SELECT
                numero_ficha,
                REPLACE(REPLACE(REPLACE(COALESCE(job, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS job,
                REPLACE(REPLACE(REPLACE(COALESCE(job_envolvidas, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS job_envolvidas,
                REPLACE(REPLACE(REPLACE(COALESCE(tipo_lamina_processo, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS tipo_lamina_processo,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_desenho_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_desenho_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(status_roteiro, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS status_roteiro,
                REPLACE(REPLACE(REPLACE(COALESCE(status_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS status_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(status_entrega_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS status_entrega_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(descricao_ferramental, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS descricao_ferramental,
                REPLACE(REPLACE(REPLACE(COALESCE(pt, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS pt,
                data_entrega_evento, data_emissao_requisicao,
                REPLACE(REPLACE(REPLACE(COALESCE(operador_responsavel, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS operador_responsavel,
                REPLACE(REPLACE(REPLACE(COALESCE(ferramenteiro_responsavel, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS ferramenteiro_responsavel,
                qtde_afiacao,
                data_termino_producao,
                REPLACE(REPLACE(REPLACE(COALESCE(observacao_ultima_manutencao_ferramenta, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS observacao_ultima_manutencao_ferramenta,
                REPLACE(REPLACE(REPLACE(COALESCE(observacao_rnc, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS observacao_rnc,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_plan_interno, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_plan_interno,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_po, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_po,
                vida_util_ferramenta,
                REPLACE(REPLACE(REPLACE(COALESCE(prensa, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS prensa,
                REPLACE(REPLACE(REPLACE(COALESCE(numero_rnc, ''), '\r', ' '), '\n', ' '), '\t', ' ') AS numero_rnc,
                total_golpes_final, total_golpes_sem_afiar,
                `1_golpes_afiacao`, `2_golpes_afiacao`, `3_golpes_afiacao`, `4_golpes_afiacao`, `5_golpes_afiacao`,
                qtde_total_afiacao, data_envio_ferramentaria, horario_inicio_ferramenteiro, horario_termino_ferramenteiro
            FROM banco_dados
            WHERE numero_ficha = %s
            """,
            (ficha_id,),
        )
        if not ficha:
            return jsonify({"error": "Ficha nao encontrada"}), 404
        return jsonify(_map_banco_dados_to_api(ficha)), 200
    except (MySQLError, RuntimeError) as error:
        return jsonify({"error": str(error)}), 400


@app.route("/api/ficha/<int:ficha_id>", methods=["PUT"])
def atualizar_ficha(ficha_id):
    try:
        return jsonify(
            {
                "success": False,
                "message": "API em modo somente leitura. Atualizacao desabilitada.",
            }
        ), 405
    except (MySQLError, RuntimeError, ValueError) as error:
        return jsonify({"success": False, "message": str(error)}), 400


@app.route("/api/ficha/<int:ficha_id>", methods=["DELETE"])
def deletar_ficha(ficha_id):
    try:
        return jsonify(
            {
                "success": False,
                "message": "API em modo somente leitura. Exclusao desabilitada.",
            }
        ), 405
    except (MySQLError, RuntimeError) as error:
        return jsonify({"success": False, "message": str(error)}), 400


@app.route("/health", methods=["GET"])
def health_check():
    try:
        row = _query_one("SELECT 1 AS ok")
        ok_value = None if not row else row.get("ok")
        if str(ok_value) == "1":
            return jsonify({"status": "ok", "storage": "mysql"}), 200
        return jsonify({"status": "error", "storage": "mysql"}), 500
    except Exception as error:
        return jsonify({"status": "error", "storage": "mysql", "message": str(error)}), 500


@app.route("/__dev_timestamp", methods=["GET"])
def dev_timestamp():
    try:
        roots = ["frontend", "."]
        exts = {".html", ".css", ".js"}
        latest = 0.0

        for root in roots:
            if not os.path.isdir(root):
                continue
            for base, _, files in os.walk(root):
                for name in files:
                    _, ext = os.path.splitext(name.lower())
                    if ext not in exts:
                        continue
                    path = os.path.join(base, name)
                    try:
                        latest = max(latest, os.path.getmtime(path))
                    except OSError:
                        pass

        return jsonify({"ts": int(latest * 1000)}), 200
    except Exception as error:
        return jsonify({"error": str(error)}), 500


if __name__ == "__main__":
    app.run(debug=True, port=5000, host="localhost")

